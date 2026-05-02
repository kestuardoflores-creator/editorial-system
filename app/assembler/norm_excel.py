import json
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

STYLE_COLS = [
    "ID_Etiqueta","Nombre_Visible","Fuente","Tamano","Interlineado",
    "Negrita","Italica","Sangria_1era","Alineacion","Color_Texto",
    "Espaciado_Antes","Espaciado_Despues","Es_Numerable","Prefijo_Texto",
    "Separador_Num","Formato_Prefijo","Formato_Numero","Posicion_Pagina","Alineacion_Objeto",
]
COL_HINTS = {
    "ID_Etiqueta":"ID (no editar)","Nombre_Visible":"Nombre visible","Fuente":"Familia tipográfica",
    "Tamano":"Tamaño en puntos","Interlineado":"1.0 / 1.5 / 2.0","Negrita":"TRUE / FALSE",
    "Italica":"TRUE / FALSE","Sangria_1era":"Puntos (0 = sin sangría)",
    "Alineacion":"LEFT / CENTER / RIGHT / JUSTIFY","Color_Texto":"Hex (#000000)",
    "Espaciado_Antes":"Puntos antes","Espaciado_Despues":"Puntos después",
    "Es_Numerable":"TRUE / FALSE","Prefijo_Texto":"Figura / Tabla / (vacío)",
    "Separador_Num":". o espacio","Formato_Prefijo":"CAPITULO_ELEMENTO / CONTINUO",
    "Formato_Numero":"ARABIC / ROMAN_UPPER / ROMAN_LOWER","Posicion_Pagina":"BREAK_TEXT / INLINE",
    "Alineacion_Objeto":"CENTER / LEFT / RIGHT",
}
COL_WIDTHS = {
    "ID_Etiqueta":18,"Nombre_Visible":24,"Fuente":18,"Tamano":10,"Interlineado":12,
    "Negrita":10,"Italica":10,"Sangria_1era":14,"Alineacion":14,"Color_Texto":13,
    "Espaciado_Antes":16,"Espaciado_Despues":17,"Es_Numerable":13,"Prefijo_Texto":14,
    "Separador_Num":12,"Formato_Prefijo":22,"Formato_Numero":18,"Posicion_Pagina":16,
    "Alineacion_Objeto":17,
}
C_HDR, C_HINT, C_LOCK, C_ALT = "1F4E79", "2E75B6", "F2F2F2", "EBF3FB"


def _fill(c): return PatternFill("solid", fgColor=c)
def _border():
    s = Side(style="thin", color="BFBFBF")
    return Border(left=s, right=s, top=s, bottom=s)
def _align(h="left"): return Alignment(horizontal=h, vertical="center", wrap_text=(h == "center"))


def _hdr_cell(ws, row, col, val, color, font_kw):
    c = ws.cell(row=row, column=col, value=val)
    c.font, c.fill, c.border = Font(name="Calibri", **font_kw), _fill(color), _border()
    c.alignment = _align("center")
    return c


def json_to_excel(json_path, excel_path):
    data = json.loads(Path(json_path).read_text(encoding="utf-8"))
    wb   = openpyxl.Workbook()
    ws   = wb.active; ws.title = "Estilos"; ws.sheet_view.showGridLines = False

    hdr_font  = dict(size=10, bold=True, color="FFFFFF")
    hint_font = dict(size=9, italic=True, color="FFFFFF")

    for i, col in enumerate(STYLE_COLS, 1):
        _hdr_cell(ws, 1, i, col, C_HDR, hdr_font)
        _hdr_cell(ws, 2, i, COL_HINTS.get(col, ""), C_HINT, hint_font)

    ws.freeze_panes = "B3"

    for ri, style in enumerate(data["estilos"], 3):
        bg = C_ALT if ri % 2 == 0 else "FFFFFF"
        for ci, col in enumerate(STYLE_COLS, 1):
            val = style.get(col, "")
            if isinstance(val, bool): val = "TRUE" if val else "FALSE"
            c = ws.cell(row=ri, column=ci, value=val)
            c.border = _border(); c.alignment = _align()
            if col == "ID_Etiqueta":
                c.fill = _fill(C_LOCK); c.font = Font(name="Calibri", size=10, bold=True, color=C_HDR)
            else:
                c.fill = _fill(bg); c.font = Font(name="Calibri", size=10)

    for i, col in enumerate(STYLE_COLS, 1):
        ws.column_dimensions[get_column_letter(i)].width = COL_WIDTHS.get(col, 14)
    ws.row_dimensions[1].height, ws.row_dimensions[2].height = 22, 18

    ws2 = wb.create_sheet("Configuracion"); ws2.sheet_view.showGridLines = False
    for i, (h, w) in enumerate([("Campo", 20), ("Valor", 20), ("Descripción", 40)], 1):
        c = ws2.cell(row=1, column=i, value=h)
        c.font, c.fill, c.border = Font(name="Calibri", size=10, bold=True, color="FFFFFF"), _fill(C_HDR), _border()
        ws2.column_dimensions[get_column_letter(i)].width = w

    m = data.get("margenes_cm", {})
    cfg_rows = [
        ("normativa",        data.get("normativa",""),       "Nombre de la normativa"),
        ("version",          data.get("version",""),         "Versión del archivo"),
        ("inicio_capitulo",  data.get("inicio_capitulo","NUEVA"), "IMPAR / PAR / NUEVA / CONTINUO"),
        ("chars_por_pagina", data.get("chars_por_pagina",2500),   "Estimado de caracteres por página"),
        ("margen_top_cm",    m.get("top",2.54),              "Margen superior en cm"),
        ("margen_bottom_cm", m.get("bottom",2.54),           "Margen inferior en cm"),
        ("margen_left_cm",   m.get("left", m.get("inner",2.54)), "Margen izquierdo en cm"),
        ("margen_right_cm",  m.get("right", m.get("outer",2.54)), "Margen derecho en cm"),
    ]
    for ri, (campo, valor, desc) in enumerate(cfg_rows, 2):
        bg = C_ALT if ri % 2 == 0 else "FFFFFF"
        for ci, (val, fkw) in enumerate([
            (campo, dict(size=10, bold=True, color=C_HDR)),
            (valor, dict(size=10)),
            (desc,  dict(size=9, italic=True, color="595959")),
        ], 1):
            c = ws2.cell(row=ri, column=ci, value=val)
            c.font = Font(name="Calibri", **fkw)
            c.fill = _fill(C_LOCK if ci == 1 else bg)
            c.border = _border()

    ws2.freeze_panes = "A2"
    wb.save(str(excel_path))


def excel_to_dict(excel_path):
    wb  = openpyxl.load_workbook(str(excel_path), data_only=True)
    ws, ws2 = wb["Estilos"], wb["Configuracion"]

    cfg = {r[0]: r[1] for r in ws2.iter_rows(min_row=2, values_only=True) if r[0]}
    headers = [c.value for c in ws[1]]
    estilos = []
    for row in ws.iter_rows(min_row=3, values_only=True):
        if not row[0]: continue
        style = {}
        for col, val in zip(headers, row):
            val = "" if val is None else val
            if str(val).upper() == "TRUE": val = True
            elif str(val).upper() == "FALSE": val = False
            style[col] = val
        estilos.append(style)

    m = {"top": cfg.get("margen_top_cm",2.54), "bottom": cfg.get("margen_bottom_cm",2.54),
         "left": cfg.get("margen_left_cm",2.54), "right": cfg.get("margen_right_cm",2.54)}
    return {"normativa": cfg.get("normativa",""), "version": cfg.get("version",""),
            "inicio_capitulo": cfg.get("inicio_capitulo","NUEVA"),
            "chars_por_pagina": cfg.get("chars_por_pagina",2500),
            "margenes_cm": m, "estilos": estilos}


def create_indice_excel(project_dir):
    """Generate config/indice.xlsx with structure template and numbering config."""
    path = Path(project_dir) / "config" / "indice.xlsx"
    wb   = openpyxl.Workbook()

    # ── Sheet 1: Estructura ───────────────────────────────────────────────────
    ws = wb.active; ws.title = "Estructura"; ws.sheet_view.showGridLines = False

    for i, (h, w, hint) in enumerate([
        ("Nivel",       8,  "1 = capítulo  2 = sección  3 = subsección …"),
        ("Titulo",      45, "Texto del encabezado tal como aparecerá en el documento"),
        ("Incluir_TOC", 13, "TRUE = aparece en la tabla de contenido  /  FALSE = no"),
        ("Notas",       30, "Comentarios opcionales (no afectan el documento)"),
    ], 1):
        _hdr_cell(ws, 1, i, h, C_HDR, dict(size=10, bold=True, color="FFFFFF"))
        _hdr_cell(ws, 2, i, hint, C_HINT, dict(size=9, italic=True, color="FFFFFF"))
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "A3"
    ws.row_dimensions[1].height, ws.row_dimensions[2].height = 22, 18

    template = [
        (1, "Introducción",               True),
        (2, "Antecedentes",               True),
        (2, "Planteamiento del Problema",  True),
        (2, "Justificación",              True),
        (2, "Objetivos",                  True),
        (1, "Marco Teórico",              True),
        (2, "Fundamentos Conceptuales",   True),
        (2, "Estado del Arte",            True),
        (1, "Metodología",               True),
        (2, "Diseño de la Investigación", True),
        (2, "Muestra y Procedimiento",    True),
        (1, "Resultados",                 True),
        (2, "Análisis de Datos",          True),
        (2, "Discusión",                  True),
        (1, "Conclusiones",               True),
    ]
    for ri, (nivel, titulo, incl) in enumerate(template, 3):
        bg   = C_ALT if ri % 2 == 0 else "FFFFFF"
        vals = [nivel, "    " * (nivel - 1) + titulo, "TRUE" if incl else "FALSE", ""]
        alns = ["center", "left", "center", "left"]
        for ci, (val, aln) in enumerate(zip(vals, alns), 1):
            c = ws.cell(row=ri, column=ci, value=val)
            c.font, c.fill, c.border = Font(name="Calibri", size=10), _fill(bg), _border()
            c.alignment = _align(aln)

    # ── Sheet 2: Numeracion ───────────────────────────────────────────────────
    ws2 = wb.create_sheet("Numeracion"); ws2.sheet_view.showGridLines = False

    for i, (h, w, hint) in enumerate([
        ("Nivel",    8,  "1–6"),
        ("Estilo",   24, "ARABIC  /  ROMAN_UPPER  /  ROMAN_LOWER  /  LETRA_UPPER  /  LETRA_LOWER  /  NINGUNA"),
        ("Separador",12, ". o espacio"),
        ("Prefijo",  20, "Capítulo / Sección / (vacío = sin prefijo)"),
    ], 1):
        _hdr_cell(ws2, 1, i, h, C_HDR, dict(size=10, bold=True, color="FFFFFF"))
        _hdr_cell(ws2, 2, i, hint, C_HINT, dict(size=9, italic=True, color="FFFFFF"))
        ws2.column_dimensions[get_column_letter(i)].width = w

    ws2.freeze_panes = "A3"
    ws2.row_dimensions[1].height, ws2.row_dimensions[2].height = 22, 18

    defaults = [
        (1, "ARABIC", ".", "Capítulo"),
        (2, "ARABIC", ".", ""),
        (3, "ARABIC", ".", ""),
        (4, "ARABIC", ".", ""),
        (5, "ARABIC", ".", ""),
        (6, "ARABIC", ".", ""),
    ]
    for ri, row in enumerate(defaults, 3):
        bg = C_ALT if ri % 2 == 0 else "FFFFFF"
        for ci, val in enumerate(row, 1):
            c = ws2.cell(row=ri, column=ci, value=val)
            c.font, c.fill, c.border = Font(name="Calibri", size=10), _fill(bg), _border()
            c.alignment = _align("center")

    wb.save(str(path))
    return path


if __name__ == "__main__":
    import sys
    if len(sys.argv) < 2: print("Uso: python norm_excel.py <normativa>"); sys.exit(1)
    base = Path(__file__).parent.parent / "config"
    jp   = base / f"{sys.argv[1]}.json"
    if not jp.exists(): print(f"❌ No encontrado: {jp}"); sys.exit(1)
    json_to_excel(jp, base / f"{sys.argv[1]}.xlsx")
    print(f"✅ Excel generado: {base / f'{sys.argv[1]}.xlsx'}")
