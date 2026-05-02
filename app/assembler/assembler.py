import argparse, json, re, sys
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_BREAK
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import openpyxl, bibtexparser

ROOT = MARKDOWNS_DIR = WORD_DIR = ASSETS_DIR = CONFIG_DIR = Path(__file__).parent.parent


def _set_project_root(project_dir):
    global ROOT, MARKDOWNS_DIR, WORD_DIR, ASSETS_DIR, CONFIG_DIR
    ROOT = Path(project_dir)
    MARKDOWNS_DIR, WORD_DIR, ASSETS_DIR, CONFIG_DIR = (
        ROOT / "markdowns", ROOT / "word", ROOT / "assets", ROOT / "config"
    )
    WORD_DIR.mkdir(exist_ok=True)


ALIGN_MAP = {
    "LEFT": WD_ALIGN_PARAGRAPH.LEFT, "CENTER": WD_ALIGN_PARAGRAPH.CENTER,
    "RIGHT": WD_ALIGN_PARAGRAPH.RIGHT, "JUSTIFY": WD_ALIGN_PARAGRAPH.JUSTIFY,
}
HEADING_MAP = {1: "TIT_CAP", 2: "H1_APA", 3: "H2_APA", 4: "H3_APA", 5: "H4_APA", 6: "H5_APA"}
CALLOUT_RE  = re.compile(r'^\s*>\s*\[!(\w+)(.*?)\]\s*$')
ATTR_RE     = re.compile(r'(\w+)="([^"]*)"')
CAPTION_RE  = re.compile(r'^\s*>\s+(.+)$')
CITE_RE     = re.compile(r'\[@([^\]]+)\]')


# ── Style engine ─────────────────────────────────────────────────────────────

def apply_style(p, s):
    pf = p.paragraph_format
    p.alignment = ALIGN_MAP.get(str(s.get("Alineacion", "LEFT")).upper(), WD_ALIGN_PARAGRAPH.LEFT)
    pf.space_before, pf.space_after = Pt(s.get("Espaciado_Antes", 0)), Pt(s.get("Espaciado_Despues", 0))
    sang = float(s.get("Sangria_1era", 0))
    if sang > 0:
        pf.first_line_indent = Pt(sang)
    elif sang < 0:
        pf.first_line_indent, pf.left_indent = Pt(sang), Pt(abs(sang))
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = float(s.get("Interlineado", 2.0))
    return p


def apply_run_style(run, s):
    color = s.get("Color_Texto", "#000000").lstrip("#")
    run.font.name  = s.get("Fuente", "Times New Roman")
    run.font.size  = Pt(float(s.get("Tamano", 12)))
    run.font.bold  = bool(s.get("Negrita", False))
    run.font.italic = bool(s.get("Italica", False))
    if len(color) == 6:
        run.font.color.rgb = RGBColor(int(color[:2], 16), int(color[2:4], 16), int(color[4:], 16))
    return run


def add_styled_paragraph(doc, text, s, citation_map=None):
    p = doc.add_paragraph()
    apply_style(p, s)
    if citation_map and CITE_RE.search(text):
        parts = CITE_RE.split(text)
        for i, chunk in enumerate(parts):
            if not chunk: continue
            t = chunk if i % 2 == 0 else _fmt_citation(chunk, citation_map)
            apply_run_style(p.add_run(t), s)
    else:
        apply_run_style(p.add_run(text), s)
    return p


def _fmt_citation(key, citation_map):
    parts = []
    for k in key.split(";"):
        k = k.strip().lstrip("@")
        e = citation_map.get(k)
        parts.append(f"{_last_name(e.get('author',''))}, {e.get('year','n.d.')}" if e else k)
    return "(" + "; ".join(parts) + ")"


def _last_name(author_field):
    if not author_field: return "?"
    a = author_field.split(" and ")[0].strip()
    return a.split(",")[0].strip() if "," in a else (a.split() or ["?"])[-1]


# ── Page breaks ───────────────────────────────────────────────────────────────

def add_page_break(doc, break_type):
    if break_type == "CONTINUO":
        return
    if break_type in ("IMPAR", "PAR"):
        t = WD_SECTION.ODD_PAGE if break_type == "IMPAR" else WD_SECTION.EVEN_PAGE
        try: doc.add_section(t); return
        except Exception: pass
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    doc.add_paragraph().add_run()._r.append(br)


# ── Numbering ─────────────────────────────────────────────────────────────────

def _to_roman(n):
    val  = [1000,900,500,400,100,90,50,40,10,9,5,4,1]
    syms = ["M","CM","D","CD","C","XC","L","XL","X","IX","V","IV","I"]
    r = ""
    for v, s in zip(val, syms):
        while n >= v: r += s; n -= v
    return r


class NumberingEngine:
    def __init__(self):
        self.chapter = 0
        self._cnt, self._gcnt = {}, {}
        self.figures, self.tables = [], []

    def next_chapter(self): self.chapter += 1

    def _fmt_n(self, n, fmt):
        fmt = str(fmt).upper() if fmt else "ARABIC"
        return _to_roman(n) if fmt == "ROMAN_UPPER" else _to_roman(n).lower() if fmt == "ROMAN_LOWER" else str(n)

    def build_prefix(self, s):
        lid, pre = s.get("ID_Etiqueta", ""), s.get("Prefijo_Texto", "")
        sep, fmt = s.get("Separador_Num", "."), s.get("Formato_Prefijo", "CONTINUO")
        fnum     = s.get("Formato_Numero", "ARABIC")
        if fmt == "CAPITULO_ELEMENTO":
            key = (lid, self.chapter)
            self._cnt[key] = self._cnt.get(key, 0) + 1
            return f"{pre} {self.chapter}{sep}{self._fmt_n(self._cnt[key], fnum)}".strip()
        self._gcnt[lid] = self._gcnt.get(lid, 0) + 1
        return f"{pre}{sep if sep.strip() else ' '}{self._fmt_n(self._gcnt[lid], fnum)}".strip()


# ── Excel table reader ────────────────────────────────────────────────────────

def read_excel_table(src, sheet):
    path = ROOT / src
    if not path.exists(): return None, None
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb.active
    data = [[str(c.value) if c.value is not None else "" for c in row] for row in ws.iter_rows()]
    return (data[0], data[1:]) if data else ([], [])


def add_excel_table(doc, headers, rows):
    if not headers: return
    tbl = doc.add_table(rows=1 + len(rows), cols=len(headers))
    tbl.style = "Table Grid"
    for j, h in enumerate(headers):
        for run in tbl.rows[0].cells[j].paragraphs[0].runs:
            run.font.bold, run.font.name, run.font.size = True, "Times New Roman", Pt(10)
        tbl.rows[0].cells[j].text = h
        tbl.rows[0].cells[j].paragraphs[0].runs[0].font.bold = True
    for i, row_data in enumerate(rows):
        for j, val in enumerate(row_data[:len(headers)]):
            tbl.rows[i + 1].cells[j].text = val


# ── BibTeX ────────────────────────────────────────────────────────────────────

def load_bib(path):
    if not Path(path).exists(): return {}
    with open(path, encoding="utf-8") as f:
        return {e["ID"]: e for e in bibtexparser.load(f).entries}


def fmt_ref_apa7(e):
    def authors(af):
        if not af: return ""
        def fmt(a):
            if "," in a:
                last, first = a.split(",", 1)
                return f"{last.strip()}, {' '.join(w[0]+'.' for w in first.split() if w)}"
            pts = a.split()
            return f"{pts[-1]}, {' '.join(p[0]+'.' for p in pts[:-1] if p)}" if pts else a
        aa = [fmt(a.strip()) for a in af.split(" and ")]
        return (", ".join(aa[:-1]) + " y " + aa[-1]) if len(aa) > 1 else aa[0]

    t, y = e.get("ENTRYTYPE", "").lower(), e.get("year", "s.f.")
    au, ti = authors(e.get("author", "")), e.get("title", "")
    doi, url = e.get("doi", ""), e.get("url", "")

    if t == "article":
        ref = f"{au} ({y}). {ti}. {e.get('journal','')}"
        if e.get("volume"): ref += f", {e['volume']}" + (f"({e['number']})" if e.get("number") else "")
        if e.get("pages"): ref += f", {e['pages'].replace('--','–')}"
        ref += "."
    elif t == "book":
        ref = f"{au} ({y}). {ti}. {e.get('publisher','')}."
    elif t in ("inproceedings", "conference"):
        ref = f"{au} ({y}). {ti}. En {e.get('booktitle','')}"
        if e.get("pages"): ref += f" (pp. {e['pages']})"
        ref += "."
    else:
        ref = f"{au} ({y}). {ti}."

    return ref + (f" https://doi.org/{doi}" if doi else f" {url}" if url else "")


# ── Page number field ─────────────────────────────────────────────────────────

def add_page_numbers(section):
    para = section.footer.paragraphs[0] if section.footer.paragraphs else section.footer.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.clear()
    run = para.add_run()
    run.font.name, run.font.size = "Times New Roman", Pt(12)
    for tag, text in [("begin", None), (None, "PAGE"), ("end", None)]:
        if tag:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), tag); run._r.append(el)
        else:
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = text; run._r.append(el)


def set_margins(section, m):
    section.top_margin    = Cm(m.get("top", 2.54))
    section.bottom_margin = Cm(m.get("bottom", 2.54))
    section.left_margin   = Cm(m.get("left", m.get("inner", 2.54)))
    section.right_margin  = Cm(m.get("right", m.get("outer", 2.54)))


# ── Markdown parser ───────────────────────────────────────────────────────────

class Elem:
    def __init__(self, kind, **kw): self.kind, self.kw = kind, kw


def parse_markdown(path):
    lines, elems, i = Path(path).read_text(encoding="utf-8").splitlines(), [], 0
    while i < len(lines):
        line = lines[i]
        if not line.strip(): i += 1; continue
        m = re.match(r'^(#{1,6})\s+(.+)$', line)
        if m:
            elems.append(Elem("heading", level=len(m.group(1)), text=m.group(2).strip()))
            i += 1; continue
        m = CALLOUT_RE.match(line)
        if m:
            tag, attrs = m.group(1).upper(), dict(ATTR_RE.findall(m.group(2)))
            i += 1; caps = []
            while i < len(lines) and CAPTION_RE.match(lines[i]):
                caps.append(CAPTION_RE.match(lines[i]).group(1)); i += 1
            elems.append(Elem("callout", tag=tag, attrs=attrs, caption=" ".join(caps))); continue
        txt = [line]; i += 1
        while i < len(lines) and lines[i].strip() and not lines[i].startswith("#") and not CALLOUT_RE.match(lines[i]):
            txt.append(lines[i]); i += 1
        elems.append(Elem("paragraph", text=" ".join(txt).strip()))
    return elems


# ── Index / TOC numbering ─────────────────────────────────────────────────────

def read_indice_config(path):
    """Read Numeracion sheet from config/indice.xlsx → {level: {estilo, sep, prefijo}}."""
    if not Path(path).exists():
        return {}
    try:
        import openpyxl as _xl
        wb = _xl.load_workbook(str(path), data_only=True)
        if "Numeracion" not in wb.sheetnames:
            return {}
        return {
            int(r[0]): {"estilo": str(r[1] or "ARABIC").upper(),
                        "sep":    str(r[2] or "."),
                        "prefijo":str(r[3] or "")}
            for r in wb["Numeracion"].iter_rows(min_row=3, values_only=True) if r[0]
        }
    except Exception:
        return {}


def heading_number(counters, level, cfg):
    """Build the TOC number string for a heading (e.g. 'Capítulo 1. ' or '1.2 ')."""
    if not cfg:
        return ""

    def fmt(n, estilo):
        if estilo == "ROMAN_UPPER":  return _to_roman(n)
        if estilo == "ROMAN_LOWER":  return _to_roman(n).lower()
        if estilo == "LETRA_UPPER":  return chr(64 + n)
        if estilo == "LETRA_LOWER":  return chr(96 + n)
        if estilo == "NINGUNA":      return ""
        return str(n)

    lc   = cfg.get(level, {})
    sep  = lc.get("sep", ".")
    pre  = lc.get("prefijo", "")
    num  = fmt(counters[level - 1], lc.get("estilo", "ARABIC"))
    if not num:
        return ""
    # Build dotted hierarchy: parent counters always use ARABIC, current uses configured style
    parents = sep.join(str(counters[l]) for l in range(level - 1) if counters[l])
    full    = f"{parents}{sep}{num}" if parents else num
    return f"{pre} {full}{sep} " if pre else f"{full} "


# ── List builders ─────────────────────────────────────────────────────────────

def build_list(doc, title, items, style_map, tit_key, ent_key):
    add_styled_paragraph(doc, title, style_map.get(tit_key, {}))
    for item in items:
        add_styled_paragraph(doc, item, style_map.get(ent_key, {}))


# ── Cover page ────────────────────────────────────────────────────────────────

def build_portada(doc, data, sm):
    s_tit, s_aut, s_inf = sm.get("PORTADA_TITULO", {}), sm.get("PORTADA_AUTOR", {}), sm.get("PORTADA_INFO", {})
    for _ in range(4): doc.add_paragraph()
    add_styled_paragraph(doc, data.get("titulo", "Título de la Tesis"), s_tit)
    doc.add_paragraph()
    add_styled_paragraph(doc, data.get("autor", "Autor"), s_aut)
    doc.add_paragraph()
    for key in ("institucion", "facultad", "programa"):
        add_styled_paragraph(doc, data.get(key, ""), s_inf)
    doc.add_paragraph()
    add_styled_paragraph(doc, f"{data.get('ciudad','')}, {data.get('anio','')}", s_inf)


# ── Main assembler ────────────────────────────────────────────────────────────

def assemble(normativa, output=None):
    xlsx, js = CONFIG_DIR / f"{normativa}.xlsx", CONFIG_DIR / f"{normativa}.json"
    if xlsx.exists():
        from norm_excel import excel_to_dict
        norm = excel_to_dict(xlsx); print(f"Normativa desde Excel: {xlsx.name}")
    elif js.exists():
        norm = json.loads(js.read_text(encoding="utf-8")); print(f"Normativa desde JSON: {js.name}")
    else:
        print(f"Normativa no encontrada: {normativa}"); sys.exit(1)

    sm       = {s["ID_Etiqueta"]: s for s in norm["estilos"]}
    ini_cap  = norm.get("inicio_capitulo", "NUEVA")
    margenes = norm.get("margenes_cm", {})
    print(f"✅ {norm['normativa']} v{norm['version']}")

    bib    = load_bib(MARKDOWNS_DIR / "referencias.bib")
    portada = {}
    port_p  = MARKDOWNS_DIR / "portada.json"
    if port_p.exists(): portada = json.loads(port_p.read_text(encoding="utf-8"))

    md_files = sorted(MARKDOWNS_DIR.glob("*.md"))
    print(f"📄 {len(md_files)} archivos Markdown")

    doc = Document()
    set_margins(doc.sections[0], margenes)
    add_page_numbers(doc.sections[0])
    for p in doc.paragraphs: p._element.getparent().remove(p._element)

    build_portada(doc, portada, sm)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    toc, figs, tabs, num, cited = [], [], [], NumberingEngine(), set()
    indice_cfg  = read_indice_config(CONFIG_DIR / "indice.xlsx")
    hcounters   = [0] * 6   # heading counters per level for TOC numbering

    for idx, md in enumerate(md_files):
        print(f"   [{idx+1}/{len(md_files)}] {md.name}")
        first = True
        for elem in parse_markdown(md):
            if elem.kind == "heading":
                lv, text = elem.kw["level"], elem.kw["text"]
                s = sm.get(HEADING_MAP.get(lv, "TEXTO_APA"), sm.get("TEXTO_APA", {}))
                if lv == 1:
                    num.next_chapter()
                    hcounters[0] += 1
                    for i in range(1, 6): hcounters[i] = 0
                    if idx > 0 or not first: add_page_break(doc, ini_cap)
                    prefix = heading_number(hcounters, 1, indice_cfg)
                    toc.append(f"{prefix}{text}")
                elif lv == 2:
                    hcounters[1] += 1
                    for i in range(2, 6): hcounters[i] = 0
                    prefix = heading_number(hcounters, 2, indice_cfg)
                    toc.append(f"    {prefix}{text}")
                elif lv >= 3:
                    hcounters[lv - 1] += 1
                    for i in range(lv, 6): hcounters[i] = 0
                    prefix = heading_number(hcounters, lv, indice_cfg)
                    toc.append(f"{'    ' * (lv - 1)}{prefix}{text}")
                add_styled_paragraph(doc, text, s, bib); first = False

            elif elem.kind == "paragraph":
                text = elem.kw["text"]
                add_styled_paragraph(doc, text, sm.get("TEXTO_APA", {}), bib)
                cited.update(k.strip().lstrip("@") for kg in CITE_RE.findall(text) for k in kg.split(";"))

            elif elem.kind == "callout":
                tag, attrs, cap = elem.kw["tag"], elem.kw["attrs"], elem.kw["caption"]
                s = sm.get(tag, sm.get("TEXTO_APA", {}))
                pre = num.build_prefix(s) if s.get("Es_Numerable") else ""

                if tag == "FIG_TIT":
                    img = ROOT / attrs.get("src", "")
                    if img.exists():
                        try:
                            ip = doc.add_paragraph(); ip.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            ip.add_run().add_picture(str(img), width=Inches(5.5))
                        except Exception as e: print(f"⚠️ Imagen: {e}")
                elif tag == "TABLA_TIT":
                    h, r = read_excel_table(attrs.get("src", ""), attrs.get("sheet"))
                    if h: add_excel_table(doc, h, r)

                if pre:
                    add_styled_paragraph(doc, pre, s)
                    if cap:
                        cs = deepcopy(s); cs["Negrita"] = False
                        add_styled_paragraph(doc, cap, cs)
                else:
                    add_styled_paragraph(doc, cap or "", s)

                if tag == "FIG_TIT": figs.append(f"{pre}  {cap}")
                elif tag == "TABLA_TIT": tabs.append(f"{pre}  {cap}")
                first = False

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    build_list(doc, "Tabla de Contenido", toc, sm, "TOC_TITULO", "TOC_ENTRADA")

    if figs:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        build_list(doc, "Lista de Figuras", figs, sm, "IDX_TITULO", "IDX_ENTRADA")
    if tabs:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        build_list(doc, "Lista de Tablas", tabs, sm, "IDX_TITULO", "IDX_ENTRADA")

    if bib:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        add_styled_paragraph(doc, "Referencias", sm.get("TOC_TITULO", {}))
        all_refs = sorted(
            [bib[k] for k in cited if k in bib] + [e for k, e in bib.items() if k not in cited],
            key=lambda e: _last_name(e.get("author", "")),
        )
        for e in all_refs:
            add_styled_paragraph(doc, fmt_ref_apa7(e), sm.get("REFERENCIA", {}))

    out = WORD_DIR / (output or "Tesis_Final.docx")
    doc.save(str(out)); print(f"\n✅ Documento: {out}")
    return out


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--normativa", default="apa7")
    ap.add_argument("--output", default=None)
    ap.add_argument("--project", default=None)
    args = ap.parse_args()
    if args.project: _set_project_root(Path(args.project).resolve())
    assemble(args.normativa, args.output)
